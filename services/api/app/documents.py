import io
import zipfile

from docx import Document as Docx
from pptx import Presentation
from pypdf import PdfReader

from . import ai, storage
from .models import Chunk, Topic


def extract(data: bytes, filename: str):
    extension = filename.rsplit(".", 1)[-1].lower()
    if extension in {"docx", "pptx"}:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Tài liệu giải nén quá lớn")
    stream = io.BytesIO(data)
    if extension == "pdf":
        if not data.startswith(b"%PDF-"):
            raise ValueError("PDF không hợp lệ")
        return [(i + 1, page.extract_text() or "") for i, page in enumerate(PdfReader(stream).pages)]
    if extension == "pptx":
        return [
            (i + 1, "\n\n".join(shape.text for shape in slide.shapes if shape.has_text_frame))
            for i, slide in enumerate(Presentation(stream).slides)
        ]
    if extension == "docx":
        doc = Docx(stream)
        return [
            (
                1,
                "\n\n".join(
                    [p.text for p in doc.paragraphs]
                    + [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
                ),
            )
        ]
    if extension == "txt":
        return [(1, data.decode("utf-8"))]
    raise ValueError("Chỉ hỗ trợ PDF, PPTX, DOCX và TXT UTF-8")


def chunk_text(text, limit=2400):
    buffer = ""
    for paragraph in text.split("\n\n"):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if len(buffer) + len(paragraph) > limit and buffer:
            yield buffer
            buffer = ""
        while len(paragraph) > limit:
            cut = paragraph.rfind(" ", 0, limit)
            cut = cut if cut > 0 else limit
            yield paragraph[:cut]
            paragraph = paragraph[cut:].strip()
        buffer += ("\n\n" if buffer else "") + paragraph
    if buffer:
        yield buffer


def process_document(db, document):
    document.embedding_model = ai.embedding_name()
    topic = db.get(Topic, document.topic_id)
    pages = extract(storage.get(document.storage_key), document.filename)
    chunks = []
    for page, text in pages:
        for content in chunk_text(text):
            chunks.append(
                Chunk(
                    document_id=document.id,
                    course_id=document.course_id,
                    topic_id=document.topic_id,
                    learning_outcome_id=topic.learning_outcome_id,
                    page=page,
                    content=content,
                    embedding=ai.embed(content),
                )
            )
            if len(chunks) > 2000:
                raise ValueError("MVP giới hạn 2000 chunks mỗi tài liệu")
    if not chunks:
        raise ValueError("Không tìm thấy văn bản; PDF scan cần OCR trước khi upload")
    db.add_all(chunks)
    document.status = "READY"
