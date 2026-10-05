"""Bounded XLSX roster parsing and formula-safe spreadsheet downloads."""
import io
import re
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .security import fail

HEADERS = ["student_number", "email", "full_name", "initial_password"]
MAX_ROWS = 5000


def workbook_bytes(headers, rows, title="Data"):
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    # Force all strings to text, including untrusted names beginning with '='.
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="175D68")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(60, max(20, len(str(column[0].value)) + 6))
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def template_bytes():
    return workbook_bytes(HEADERS, [["SV001", "sv001@example.edu.vn", "Nguyễn Văn An", ""]], "Students")


def parse_roster(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 200 or sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                fail(413, "EXCEL_SIZE", "File Excel giải nén quá lớn")
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = book.active
            if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > 4:
                fail(422, "EXCEL_DIMENSIONS", "Tối đa 5000 sinh viên và đúng 4 cột theo mẫu")
            rows = sheet.iter_rows(max_row=min(sheet.max_row or MAX_ROWS + 2, MAX_ROWS + 2), max_col=5)
            header = next(rows)
            if [c.value for c in header[:4]] != HEADERS or header[4].value is not None:
                fail(422, "EXCEL_HEADER", "Dùng đúng 4 cột trong file Excel mẫu")
            result, emails, numbers = [], set(), set()
            for index, cells in enumerate(rows, 2):
                if all(c.value is None for c in cells):
                    continue
                if index > MAX_ROWS + 1:
                    fail(422, "EXCEL_ROWS", "Tối đa 5000 dòng sinh viên")
                if any(c.data_type == "f" for c in cells):
                    fail(422, "EXCEL_FORMULA", f"Dòng {index}: không chấp nhận công thức")
                number, email, name, password, extra = [str(c.value).strip() if c.value is not None else "" for c in cells]
                email = email.lower()
                if (extra or not number or len(number) > 80 or not name or len(name) > 150
                        or len(email) > 320 or not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email)
                        or (password and not 12 <= len(password) <= 128)):
                    fail(422, "EXCEL_ROW", f"Dòng {index}: kiểm tra MSSV, email, họ tên và mật khẩu (12–128 ký tự nếu nhập)")
                if email in emails or number in numbers:
                    fail(422, "EXCEL_DUPLICATE", f"Dòng {index}: trùng email hoặc MSSV")
                emails.add(email)
                numbers.add(number)
                result.append({"student_number": number, "email": email, "name": name, "password": password})
            if not result:
                fail(422, "EXCEL_EMPTY", "Danh sách sinh viên trống")
            return result
        finally:
            book.close()
    except (zipfile.BadZipFile, KeyError, ValueError, OSError, StopIteration) as exc:
        fail(422, "INVALID_EXCEL", f"Không đọc được file .xlsx ({type(exc).__name__})")
