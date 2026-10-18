"""Atomic Excel imports with a non-writing validation preview and bounded input."""
import io
import re
import zipfile
from collections import OrderedDict

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .db import get_db
from .excel import workbook_bytes
from .models import Audit, Course, Exam, Rubric, Topic, User
from .schemas import CourseIn, ExamIn, QuestionOutput, RubricIn, UserIn
from .security import course_access, fail, hasher, staff

router = APIRouter(prefix="/admin/imports")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SPECS = {
    "courses": (["course_code", "course_name", "description", "teacher_email"],
        [["ENG101", "Tiếng Anh vấn đáp", "Môn học tiếng Anh", "teacher@example.edu.vn"]]),
    "rubrics": (["course_code", "rubric_name", "criterion_name", "description", "max_score", "weight"],
        [["ENG101", "Rubric vấn đáp", "Nội dung", "Trả lời đúng và đủ ý", 10, 6],
         ["ENG101", "Rubric vấn đáp", "Diễn đạt", "Mạch lạc và có ví dụ", 10, 4]]),
    "exams": (["course_code", "exam_name", "rubric_name", "time_minutes", "question_number", "topic_name", "difficulty", "question_text", "expected_concepts", "english_terms"],
        [["ENG101", "Vấn đáp tiếng Anh", "Rubric vấn đáp", 10, 1, "Daily life", "EASY", "Describe your daily study routine and give an example.", "study routine; specific example", "routine=thói quen"],
         ["ENG101", "Vấn đáp tiếng Anh", "Rubric vấn đáp", 10, 2, "Daily life", "MEDIUM", "Describe a challenge you faced and how you solved it.", "challenge; solution", "challenge=thử thách"]]),
    "teachers": (["username", "email", "full_name", "initial_password"],
        [["gv001", "teacher@example.edu.vn", "Nguyễn Văn An", ""]]),
    "officers": (["username", "email", "full_name", "initial_password"],
        [["kt001", "officer@example.edu.vn", "Trần Thị Bình", ""]]),
}
NOTES = {
    "courses": "Mỗi dòng một môn. Admin có thể để trống teacher_email để tự phụ trách. Khảo thí cần email giảng viên đã nhập. Giáo viên chỉ tạo môn do mình phụ trách.",
    "rubrics": "Một dòng cho mỗi tiêu chí. Các dòng cùng course_code + rubric_name tạo một rubric (tối đa 20 tiêu chí). max_score và weight là số dương, tối đa 100.",
    "exams": "Một dòng một câu hỏi có sẵn. Cùng course_code + exam_name: rubric_name, time_minutes phải giống nhau; question_number liên tiếp từ 1, tối đa 20. topic_name và rubric_name phải tồn tại duy nhất trong môn. difficulty: EASY/MEDIUM/HARD. expected_concepts ngăn cách bằng ;. english_terms tùy chọn: thuật ngữ=nghĩa; thuật ngữ=nghĩa. Đề được lưu DRAFT, không gọi AI tạo câu hỏi/rubric, không tự công bố. Chuẩn bị tài liệu READY theo chủ đề rồi tạo bản nháp để gắn nguồn và duyệt trong Điều phối kỳ thi.",
    "teachers": "Chỉ admin nhập. Tài khoản tạo mới có role TEACHER. Nhập mật khẩu ban đầu riêng (12–128 ký tự) và phân phối qua kênh của đơn vị. Không đổi role/mật khẩu tài khoản đã tồn tại; không tự liên kết Google cho tài khoản đặc quyền.",
    "officers": "Chỉ admin nhập. Tài khoản tạo mới có role EXAM_OFFICER. Nhập mật khẩu ban đầu riêng (12–128 ký tự). Không đổi role/mật khẩu tài khoản đã tồn tại; không tự liên kết Google cho tài khoản đặc quyền.",
}


def access(kind, user):
    if kind not in SPECS:
        fail(404, "IMPORT_KIND", "Loại import không hỗ trợ")
    if kind in {"teachers", "officers"} and user.role != "ADMIN":
        fail(403, "FORBIDDEN", "Chỉ admin được import giảng viên và khảo thí")


@router.get("/{kind}/template.xlsx")
def template(kind: str, course_id: str | None = None, db: Session = Depends(get_db), user=Depends(staff)):
    access(kind, user)
    headers, examples = SPECS[kind]
    examples = [list(r) for r in examples]
    refs = []
    if course_id and kind in {"rubrics", "exams"}:
        course = course_access(db, course_id, user)
        for r in examples:
            r[0] = course.code
        if kind == "exams":
            rubrics = db.scalars(select(Rubric).where(Rubric.course_id == course_id).order_by(Rubric.name)).all()
            topics = db.scalars(select(Topic).where(Topic.course_id == course_id).order_by(Topic.name)).all()
            for r in examples:
                if rubrics:
                    r[2] = rubrics[0].name
                if topics:
                    r[5] = topics[0].name
            refs = [["rubric_name", r.name] for r in rubrics] + [["topic_name", t.name] for t in topics]
    if kind == "courses" and user.role == "TEACHER":
        examples[0][3] = user.email or ""
    book = load_workbook(io.BytesIO(workbook_bytes(headers, examples, "Data")))
    guide = book.create_sheet("Huong dan")
    guide.append(["Hướng dẫn"])
    guide.append([NOTES[kind]])
    guide.append(["Chỉ nhập sheet Data; các dòng mẫu sẽ được nhập nếu không thay/xóa. Tối đa 2 MB, 1.000 dòng (tài khoản: 200). Không dùng công thức. File có lỗi sẽ không ghi bất kỳ dòng nào."])
    guide.append(["Import không ghi đè: môn/rubric/tài khoản giống hệt được bỏ qua; trùng khóa khác nội dung báo lỗi. Tên đề đã tồn tại báo lỗi, hãy đổi tên để tạo đề mới."])
    if refs:
        sheet = book.create_sheet("Danh muc")
        for row in [["field", "value"], *refs]:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                cell.data_type = "s"
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return Response(stream.getvalue(), media_type=XLSX,
        headers={"Content-Disposition": f'attachment; filename="{kind}-template.xlsx"'})


def parse(file, kind, custom_headers=None):
    if not (file.filename or "").lower().endswith(".xlsx"):
        fail(422, "EXCEL_TYPE", "Chọn file .xlsx theo mẫu")
    raw = file.file.read(2 * 1024 * 1024 + 1)
    if len(raw) > 2 * 1024 * 1024:
        fail(413, "EXCEL_SIZE", "File Excel tối đa 2 MB")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 200 or sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                fail(413, "EXCEL_SIZE", "File giải nén quá lớn")
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = book["Data"] if "Data" in book.sheetnames else book.active
            headers = custom_headers if custom_headers is not None else SPECS[kind][0]
            limit = 200 if kind in {"teachers", "officers"} else 1000
            if (sheet.max_row or 0) > limit + 1 or (sheet.max_column or 0) > len(headers):
                fail(422, "EXCEL_DIMENSIONS", f"Tối đa {limit} dòng và đúng số cột trong mẫu")
            rows = sheet.iter_rows(max_row=min(sheet.max_row or limit + 2, limit + 2), max_col=len(headers) + 1)
            first = next(rows)
            if [c.value for c in first] != headers + [None]:
                fail(422, "EXCEL_HEADER", "Giữ nguyên tên và thứ tự các cột trong mẫu")
            result = []
            for index, cells in enumerate(rows, 2):
                if all(c.value is None for c in cells):
                    continue
                if index > limit + 1 or cells[-1].value is not None:
                    fail(422, "EXCEL_DIMENSIONS", "Quá số dòng/cột cho phép")
                if any(c.data_type in {"f", "e"} for c in cells):
                    fail(422, "EXCEL_FORMULA", f"Dòng {index}: không dùng công thức hoặc ô lỗi")
                values = [str(c.value).strip() if c.value is not None else "" for c in cells[:-1]]
                if any(len(v) > 10000 for v in values):
                    fail(422, "EXCEL_CELL", f"Dòng {index}: nội dung ô quá dài")
                result.append((index, dict(zip(headers, values))))
            if not result:
                fail(422, "EXCEL_EMPTY", "File không có dữ liệu")
            return result
        finally:
            book.close()
    except HTTPException:
        raise
    except Exception:
        fail(422, "INVALID_EXCEL", "Không đọc được file Excel. Lưu lại dưới dạng .xlsx rồi thử lại.")


def unique(db, model, predicate, label):
    rows = db.scalars(select(model).where(predicate)).all()
    if len(rows) != 1:
        raise ValueError(f"{label} phải tồn tại duy nhất; kiểm tra tên/mã/email trong hệ thống")
    return rows[0]


def course_for(db, code, user):
    c = unique(db, Course, func.lower(Course.code) == code.lower(), "Mã môn")
    course_access(db, c.id, user)
    if c.status != "ACTIVE":
        raise ValueError("Môn đã lưu trữ")
    return c


def build_plan(db, kind, rows, user):
    plan, errors, groups, seen = [], [], OrderedDict(), set()
    def error(line, exc):
        if isinstance(exc, ValidationError):
            message = "; ".join(f"{'.'.join(map(str, e['loc']))}: {e['msg']}" for e in exc.errors(include_input=False)[:5])
        else:
            message = str(exc)
        if len(errors) < 100:
            errors.append({"row": line, "message": message})
    for line, row in rows:
        try:
            if kind in {"rubrics", "exams"}:
                key = (row["course_code"].lower(), row["rubric_name"] if kind == "rubrics" else row["exam_name"])
                groups.setdefault(key, []).append((line, row))
                continue
            if kind == "courses":
                body = CourseIn(code=row["course_code"], name=row["course_name"], description=row["description"])
                if body.code.lower() in seen:
                    raise ValueError("Mã môn bị lặp trong file")
                seen.add(body.code.lower())
                owner = user
                if row["teacher_email"]:
                    owner = unique(db, User, func.lower(User.email) == row["teacher_email"].lower(), "Email giảng viên")
                    if owner.role != "TEACHER" or owner.status != "ACTIVE":
                        raise ValueError("Người phụ trách phải là giảng viên đang hoạt động")
                elif user.role == "EXAM_OFFICER":
                    raise ValueError("Khảo thí cần nhập email giảng viên phụ trách")
                if user.role == "TEACHER" and owner.id != user.id:
                    raise ValueError("Giáo viên chỉ nhập môn do mình phụ trách")
                data = body.model_dump() | {"owner_id": owner.id}
                old = db.scalar(select(Course).where(func.lower(Course.code) == body.code.lower()))
                if old and (old.status != "ACTIVE" or any(getattr(old, k) != v for k, v in data.items())):
                    raise ValueError("Mã môn đã tồn tại với nội dung/người phụ trách khác; không ghi đè")
                plan.append({"action": "skip" if old else "create", "label": body.code, "data": data})
            else:
                role = "TEACHER" if kind == "teachers" else "EXAM_OFFICER"
                email = row["email"].lower()
                if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
                    raise ValueError("Email không hợp lệ")
                if ("email", email) in seen or ("username", row["username"].lower()) in seen:
                    raise ValueError("Trùng email hoặc tên đăng nhập trong file")
                seen.update({("email", email), ("username", row["username"].lower())})
                existing = db.scalars(select(User).where((func.lower(User.email) == email) | (func.lower(User.username) == row["username"].lower()))).all()
                if existing:
                    if len(existing) != 1 or existing[0].role != role or existing[0].status != "ACTIVE" or existing[0].username != row["username"] or (existing[0].email or "").lower() != email or existing[0].name != row["full_name"]:
                        raise ValueError("Email/tên đăng nhập đã tồn tại với danh tính hoặc vai trò khác; không thay đổi tài khoản")
                    plan.append({"action": "skip", "label": email, "data": {}})
                else:
                    data = UserIn(username=row["username"], email=email, name=row["full_name"], role=role, password=row["initial_password"]).model_dump()
                    plan.append({"action": "create", "label": email, "data": data})
        except (ValueError, ValidationError) as exc:
            error(line, exc)
    for _, items in groups.items():
        line, first = items[0]
        try:
            course = course_for(db, first["course_code"], user)
            if kind == "rubrics":
                criteria = []
                for line, r in items:
                    criteria.append({"name": r["criterion_name"], "description": r["description"], "max_score": r["max_score"], "weight": r["weight"]})
                body = RubricIn(name=first["rubric_name"], criteria=criteria)
                data = body.model_dump() | {"course_id": course.id}
                existing = db.scalars(select(Rubric).where(Rubric.course_id == course.id, Rubric.name == body.name)).all()
                if existing and (len(existing) != 1 or existing[0].criteria != data["criteria"]):
                    raise ValueError("Rubric trùng tên với nội dung khác; đổi tên để tạo rubric mới")
                plan.append({"action": "skip" if existing else "create", "label": f"{course.code} · {body.name}", "data": data})
            else:
                rubric = unique(db, Rubric, (Rubric.course_id == course.id) & (Rubric.name == first["rubric_name"]), "Rubric")
                if db.scalar(select(Exam.id).where(Exam.course_id == course.id, Exam.name == first["exam_name"])):
                    raise ValueError("Tên bài thi đã tồn tại trong môn; đổi tên để tạo bài mới")
                ordered = sorted(items, key=lambda pair: int(pair[1]["question_number"]))
                if [int(r["question_number"]) for _, r in ordered] != list(range(1, len(items) + 1)):
                    raise ValueError("Số câu phải liên tiếp từ 1, không trùng hoặc thiếu")
                questions, blueprint = [], []
                for line, r in ordered:
                    if r["rubric_name"] != first["rubric_name"] or int(r["time_minutes"]) != int(first["time_minutes"]):
                        raise ValueError("Các dòng của một đề phải cùng rubric và thời lượng")
                    topic = unique(db, Topic, (Topic.course_id == course.id) & (Topic.name == r["topic_name"]), "Chủ đề")
                    terms = []
                    for term in filter(None, (x.strip() for x in r["english_terms"].split(";"))):
                        key, sep, meaning = term.partition("=")
                        if not sep:
                            raise ValueError("english_terms dùng dạng term=meaning;term=meaning")
                        terms.append({"term": key.strip(), "meaning": meaning.strip()})
                    question = QuestionOutput(text=r["question_text"], expected_concepts=[x.strip() for x in r["expected_concepts"].split(";") if x.strip()], english_terms=terms, reference_chunk_ids=["pending-source"])
                    questions.append(question.model_dump() | {"reference_chunk_ids": [], "topic_id": topic.id, "difficulty": r["difficulty"]})
                    blueprint.append({"topic_id": topic.id, "difficulty": r["difficulty"], "count": 1})
                body = ExamIn(course_id=course.id, rubric_id=rubric.id, name=first["exam_name"], time_limit=int(first["time_minutes"]) * 60, blueprint=blueprint)
                plan.append({"action": "create", "label": f"{course.code} · {body.name} ({len(questions)} câu)", "data": body.model_dump() | {"workflow": True, "imported_questions": questions}})
        except (ValueError, ValidationError) as exc:
            error(line, exc)
    return plan, errors


def outcome(rows, plan, errors):
    return {"valid": not errors, "rows": len(rows), "create": sum(p["action"] == "create" for p in plan),
            "skip": sum(p["action"] == "skip" for p in plan), "errors": errors,
            "items": [{"label": p["label"], "action": p["action"]} for p in plan]}


@router.post("/{kind}/preview")
def preview(kind: str, file: UploadFile = File(), db: Session = Depends(get_db), user=Depends(staff)):
    access(kind, user)
    rows = parse(file, kind)
    plan, errors = build_plan(db, kind, rows, user)
    return outcome(rows, plan, errors)


@router.post("/{kind}")
def apply(kind: str, file: UploadFile = File(), db: Session = Depends(get_db), user=Depends(staff)):
    access(kind, user)
    rows = parse(file, kind)
    # Serializes import-vs-import uniqueness checks, including rubric/exam names
    # without a legacy database uniqueness constraint.
    db.scalar(select(User).where(User.role == "ADMIN", User.status == "ACTIVE").order_by(User.id).limit(1).with_for_update())
    plan, errors = build_plan(db, kind, rows, user)
    if errors:
        fail(422, "IMPORT_INVALID", "; ".join(f"Dòng {e['row']}: {e['message']}" for e in errors[:5]))
    model = {"courses": Course, "rubrics": Rubric, "exams": Exam, "teachers": User, "officers": User}[kind]
    for item in plan:
        if item["action"] == "skip":
            continue
        data = item["data"].copy()
        if model == User:
            data["password_hash"] = hasher.hash(data.pop("password"))
        db.add(model(**data))
    db.add(Audit(user_id=user.id, event="EXCEL_IMPORTED", details={"kind": kind, "rows": len(rows), "created": sum(p["action"] == "create" for p in plan)}))
    db.commit()
    return outcome(rows, plan, [])
