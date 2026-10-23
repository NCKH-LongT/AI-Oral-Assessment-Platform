import io
import time
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from test_mvp import ok, prepare

from app import speech, worker
from app.excel import HEADERS, workbook_bytes
from app.models import Assignment, Attempt, CourseEnrollment, Exam, ExamSession, SystemSetting, User


def setup_workflow(env):
    context = prepare(env, 1)
    clients, factory = env
    admin, officer, teacher = clients["admin"], clients["reviewer"], clients["teacher"]
    uid = ok(officer.get("/auth/me"))["id"]
    ok(admin.put(f"/admin/users/{uid}/role", json={"role": "EXAM_OFFICER"}))
    tid = ok(teacher.get("/auth/me"))["id"]
    exam = ok(officer.post("/admin/examination/exams", json={
        "course_id": context["course"]["id"], "teacher_id": tid, "name": "Final examination", "time_limit": 900,
    }), 201)
    key = exam["id"]
    with factory() as db:
        rubric_id = db.get(Exam, key).rubric_id
        db.scalar(select(User).where(User.username == "student")).email = "student@example.edu.vn"
        db.commit()
    payload = context["payload"] | {"rubric_id": rubric_id, "name": "Final examination"}
    ok(teacher.put(f"/admin/exams/{key}", json=payload))
    return clients, factory, key, context


def import_roster(client, key, rows=None):
    raw = workbook_bytes(HEADERS, rows or [["0001", "student@example.edu.vn", "Student", ""]])
    return client.post(f"/admin/exams/{key}/import-students", files={"file": ("students.xlsx", raw)})


def test_select_course_and_class_roster(env):
    clients, factory, key, context = setup_workflow(env)
    office, admin = clients["reviewer"], clients["admin"]
    course_id = context["course"]["id"]
    sid = ok(clients["student"].get("/auth/me"))["id"]
    other = ok(clients["outsider"].get("/auth/me"))["id"]
    ok(admin.post(f"/admin/courses/{course_id}/students", json={"student_ids": [sid, other]}))
    section = ok(admin.post("/admin/classes", json={"code": "SE1801", "name": "SE1801", "course_id": course_id, "starts_at": 1, "ends_at": 9999999999}), 201)
    ok(admin.post(f"/admin/classes/{section['id']}/students", json={"student_ids": [sid]}))
    ok(admin.delete(f"/admin/courses/{course_id}/students/{sid}"))
    assert ok(office.get(f"/admin/courses/{course_id}/classes"))[0]["student_count"] == 1
    endpoint = f"/admin/exams/{key}/roster-candidates"
    students = ok(office.get(endpoint, params={"course_id": course_id}))
    assert {s["id"] for s in students} == {sid, other}
    params = {"course_id": course_id, "class_id": section["id"]}
    assert [s["id"] for s in ok(office.get(endpoint, params=params))] == [sid]
    body = params | {"student_ids": [sid, other]}
    assert office.post(f"/admin/exams/{key}/select-students", json=body).status_code == 422
    assert ok(office.get(f"/admin/exams/{key}/roster")) == []
    body["student_ids"] = [sid, sid]
    assert ok(office.post(f"/admin/exams/{key}/select-students", json=body)) == {"added": 1, "skipped": 0}
    assert ok(office.get(endpoint, params=params))[0]["assigned"] is True
    ok(import_roster(office, key))  # Excel can supply MSSV for an existing selection.
    with factory() as db:
        row = db.scalar(select(Assignment).where(Assignment.exam_id == key))
        row.opens_at, row.closes_at = 100, 200
        db.commit()
    assert ok(office.post(f"/admin/exams/{key}/select-students", json=body)) == {"added": 0, "skipped": 1}
    row = ok(office.get(f"/admin/exams/{key}/roster"))[0]
    assert (row["student_number"], row["opens_at"], row["closes_at"]) == ("0001", 100, 200)
    body = {"course_id": course_id, "student_ids": [other]}
    assert ok(office.post(f"/admin/exams/{key}/select-students", json=body))["added"] == 1
    for role in ("teacher", "student"):
        assert clients[role].get(endpoint, params=params).status_code == 403
        assert clients[role].post(f"/admin/exams/{key}/select-students", json=body).status_code == 403
    with factory() as db:
        db.get(Exam, key).status = "PUBLISHED"
        db.commit()
    assert office.post(f"/admin/exams/{key}/select-students", json=body).status_code == 409


@pytest.mark.parametrize("membership", ["course", "class"])
def test_roster_validates_course_and_active_membership(env, membership):
    clients, factory, key, context = setup_workflow(env)
    office, admin = clients["reviewer"], clients["admin"]
    course_id = context["course"]["id"]
    student_id = ok(clients["student"].get("/auth/me"))["id"]
    section = ok(admin.post("/admin/classes", json={
        "code": "SE1801", "name": "SE1801", "course_id": course_id,
        "starts_at": 1, "ends_at": 9999999999,
    }), 201)
    source = {"course_id": course_id}
    if membership == "class":
        source["class_id"] = section["id"]
    candidates = f"/admin/exams/{key}/roster-candidates"
    selection = f"/admin/exams/{key}/select-students"
    body = source | {"student_ids": [student_id]}
    assert office.post(selection, json=body).status_code == 422
    membership_url = (f"/admin/classes/{section['id']}/students" if membership == "class"
                      else f"/admin/courses/{course_id}/students")
    ok(admin.post(membership_url, json={"student_ids": [student_id]}))
    assert [student["id"] for student in ok(office.get(candidates, params=source))] == [student_id]
    other_course = ok(admin.post("/admin/courses", json={"code": "OTHER", "name": "Other", "description": ""}), 201)
    wrong_source = {"course_id": other_course["id"], "class_id": section["id"]}
    assert office.get(candidates, params=wrong_source).status_code == 422
    assert office.post(selection, json=wrong_source | {"student_ids": [student_id]}).status_code == 422
    with factory() as db:
        db.get(User, student_id).status = "DISABLED"
        db.commit()
    assert ok(office.get(candidates, params=source)) == []
    assert office.post(selection, json=body).status_code == 422


def approve_and_schedule(clients, key, opens=None):
    teacher, office = clients["teacher"], clients["reviewer"]
    ok(teacher.post(f"/admin/exams/{key}/generate"))
    ok(teacher.post(f"/admin/exams/{key}/teacher-approve"))
    sid = ok(clients["student"].get("/auth/me"))["id"]
    ok(office.put(f"/admin/exams/{key}/schedule", json={
        "student_ids": [sid], "opens_at": opens or time.time() - 1, "closes_at": time.time() + 3600,
    }))
    ok(office.post(f"/admin/exams/{key}/office-approve"))


def test_workflow_permissions_review_and_schedule(env):
    clients, factory, key, context = setup_workflow(env)
    teacher, office, student, outsider = [clients[k] for k in ("teacher", "reviewer", "student", "outsider")]
    assert import_roster(teacher, key).status_code == 403
    assert ok(import_roster(office, key))["added"] == 1
    assert ok(import_roster(office, key))["added"] == 0
    assert teacher.post(f"/admin/exams/{key}/publish").status_code == 409
    assert office.post(f"/admin/exams/{key}/generate").status_code == 403
    assert office.post(f"/admin/exams/{key}/office-approve").status_code == 409
    generated = ok(teacher.post(f"/admin/exams/{key}/generate"))
    draft = generated["snapshot"]
    assert generated["status"] == "GENERATED"
    question = {k: draft["questions"][0][k] for k in ("text", "expected_concepts", "reference_chunk_ids", "english_terms")}
    question["text"] = "Giải thích Dependency Injection và đưa ví dụ minh họa."
    ok(teacher.put(f"/admin/exams/{key}/draft", json={"questions": [question], "criteria": draft["criteria"]}))
    assert student.get(f"/admin/exams/{key}/draft").status_code == 403
    assert student.post("/exam-sessions", json={"exam_id": key}).status_code == 403
    ok(teacher.post(f"/admin/exams/{key}/teacher-approve"))
    assert teacher.put(f"/admin/exams/{key}/draft", json={"questions": [question], "criteria": draft["criteria"]}).status_code == 409
    assert office.post(f"/admin/exams/{key}/office-approve").status_code == 409
    sid = ok(student.get("/auth/me"))["id"]
    future = time.time() + 1000
    schedule = {"student_ids": [sid], "opens_at": future, "closes_at": future + 3600}
    assert teacher.put(f"/admin/exams/{key}/schedule", json=schedule).status_code == 403
    ok(office.put(f"/admin/exams/{key}/schedule", json=schedule))
    ok(office.post(f"/admin/exams/{key}/office-approve"))
    available = ok(student.get("/exams/available"))
    assert next(e for e in available if e["id"] == key)["opens_at"] == future
    assert student.post("/exam-sessions", json={"exam_id": key}).status_code == 409
    with factory() as db:
        a = db.scalar(select(Assignment).where(Assignment.exam_id == key))
        a.opens_at = time.time() - 1
        other_id = db.scalar(select(User.id).where(User.username == "outsider"))
        db.add(CourseEnrollment(course_id=context["course"]["id"], student_id=other_id))
        db.commit()
    assert key not in [e["id"] for e in ok(outsider.get("/exams/available"))]
    assert outsider.post("/exam-sessions", json={"exam_id": key}).status_code == 403
    session = ok(student.post("/exam-sessions", json={"exam_id": key}))
    with factory() as db:
        db.scalar(select(Assignment).where(Assignment.exam_id == key)).closes_at = time.time() - 1
        db.commit()
    assert student.post(f"/exam-sessions/{session['id']}/start").status_code == 409
    assert import_roster(office, key).status_code == 409


def test_excel_atomicity_email_login_and_export(env):
    clients, factory, key, _ = setup_workflow(env)
    office = clients["reviewer"]
    rows = [["0002", "new@example.edu.vn", "New student", "initial-password-123"],
            ["0003", "conflict@example.edu.vn", "Conflict", ""]]
    with factory() as db:
        db.scalar(select(User).where(User.username == "teacher")).email = "conflict@example.edu.vn"
        db.commit()
    assert import_roster(office, key, rows).status_code == 422
    with factory() as db:
        assert db.scalar(select(User).where(User.email == "new@example.edu.vn")) is None
        assert db.scalar(select(Assignment).where(Assignment.exam_id == key)) is None
    assert ok(import_roster(office, key, rows[:1]))["accounts_created"] == 1
    student = clients["student"]
    user = ok(student.post("/auth/login", json={"username": "NEW@example.edu.vn", "password": "initial-password-123"}))["user"]
    assert user["email"] == "new@example.edu.vn"
    assert import_roster(office, key, rows[:1] * 2).status_code == 422
    result = office.get(f"/admin/exams/{key}/grades.xlsx")
    assert result.status_code == 200
    book = load_workbook(io.BytesIO(result.content))
    assert book.active.cell(2, 1).value == "0002"
    assert book.active.cell(2, 5).value == "NOT_STARTED"
    assert book.active.cell(2, 7).value is None
    assert student.get(f"/admin/exams/{key}/grades.xlsx").status_code == 403


@pytest.mark.parametrize("score,expected", [(0, "REVIEW_REQUIRED"), (4.99, "REVIEW_REQUIRED"), (5, "COMPLETED"), (8, "COMPLETED")])
def test_low_grade_manual_override_and_export(env, score, expected):
    clients, factory, key, _ = setup_workflow(env)
    office, teacher, student = clients["reviewer"], clients["teacher"], clients["student"]
    ok(import_roster(office, key))
    approve_and_schedule(clients, key)
    session = ok(student.post("/exam-sessions", json={"exam_id": key}))
    with factory() as db:
        row = db.get(ExamSession, session["id"])
        row.status = "SUBMITTED"
        attempt = db.scalar(select(Attempt).where(Attempt.session_id == row.id))
        attempt.status, attempt.assessment = "GRADED", {"score": score, "review_required": False}
        db.flush()
        worker.finalize(db, row)
        assert row.ai_score == score and row.status == expected
        assert row.final_score == (score if expected == "COMPLETED" else None)
        db.commit()
    payload = {"score": 6.5, "reason": "Đã kiểm tra bản ghi và nội dung câu trả lời."}
    assert office.post(f"/admin/results/{session['id']}/manual-grade", json=payload).status_code == 403
    ok(teacher.post(f"/admin/results/{session['id']}/manual-grade", json=payload))
    assert "final_score" not in ok(student.get(f"/exam-sessions/{session['id']}"))
    result = ok(teacher.get(f"/admin/results/{session['id']}"))
    assert result["status"] == "COMPLETED" and result["final_score"] == 6.5
    with factory() as db:
        row = db.get(ExamSession, session["id"])
        worker.finalize(db, row)
        assert row.final_score == 6.5 and row.ai_score == score
    book = load_workbook(io.BytesIO(office.get(f"/admin/exams/{key}/grades.xlsx").content))
    assert book.active.cell(2, 6).value == score
    assert book.active.cell(2, 7).value == 6.5


def test_request_role_requires_admin_approval(env):
    clients, _ = env
    student = clients["student"]
    ok(student.post("/auth/request-exam-officer"))
    user = ok(student.get("/auth/me"))
    assert user["role"] == "STUDENT" and user["requested_role"] == "EXAM_OFFICER"
    assert student.get("/admin/examination/teachers").status_code == 403
    ok(clients["admin"].put(f"/admin/users/{user['id']}/role", json={"role": "EXAM_OFFICER"}))
    assert ok(student.get("/auth/me"))["requested_role"] is None
    assert student.get("/admin/examination/teachers").status_code == 200


def test_hotwords_are_frozen_scoped_and_forwarded(env, monkeypatch):
    clients, factory, key, context = setup_workflow(env)
    teacher, office, student = clients["teacher"], clients["reviewer"], clients["student"]
    path = f"/admin/courses/{context['course']['id']}/hotwords"
    words = [f"Thuật ngữ chuyên ngành {i}" for i in range(200)]
    ok(teacher.put(path, json={"hotwords": words}))
    # Old global lists, including those larger than the former limit, are ignored.
    with factory() as db:
        db.add(SystemSetting(key="speech", value={"provider": "local_server", "hotwords": ["Global"] * 200}))
        db.commit()
    assert "hotwords" not in ok(student.get("/stt/config"))
    assert "hotwords" not in ok(clients["admin"].get("/admin/settings/speech"))
    ok(import_roster(office, key))
    approve_and_schedule(clients, key)
    ok(teacher.put(path, json={"hotwords": ["Changed vocabulary"]}))
    session = ok(student.post("/exam-sessions", json={"exam_id": key}))
    config = ok(student.get("/stt/config", params={"session_id": session["id"]}))
    assert config["hotwords"] == words
    assert clients["outsider"].get("/stt/config", params={"session_id": session["id"]}).status_code == 403
    seen = {}
    def recognize(path, policy):
        seen.update(policy)
        return {"transcript": "PostgreSQL", "stt_confidence": .9}
    monkeypatch.setattr("app.stt.transcribe_file", recognize)
    ok(student.post("/stt", data={"session_id": session["id"]}, files={"file": ("a.webm", b"audio")}))
    assert seen["hotwords"] == config["hotwords"]
    def transcribe(*args, **kwargs):
        seen.update(kwargs)
        return [SimpleNamespace(text="PostgreSQL", avg_logprob=-.1)], SimpleNamespace(language="vi")
    monkeypatch.setattr(speech, "model", lambda _: SimpleNamespace(transcribe=transcribe))
    speech.whisper("audio.wav", "vi", config["hotwords"])
    assert seen["hotwords"] == ", ".join(words)


def test_course_hotwords_are_independent_and_validate_before_saving(env):
    clients, factory = env
    admin, teacher, student = clients["admin"], clients["teacher"], clients["student"]
    first = ok(teacher.post("/admin/courses", json={"code": "CS", "name": "Computing"}), 201)
    second = ok(admin.post("/admin/courses", json={"code": "BIO", "name": "Biology"}), 201)
    first_path, second_path = [f"/admin/courses/{course['id']}/hotwords" for course in (first, second)]
    words = [str(i).ljust(20, "x") for i in range(500)]
    assert ok(teacher.put(first_path, json={"hotwords": words + ["", " ", " " + words[0] + " "]}))["hotwords"] == words
    ok(admin.put(second_path, json={"hotwords": ["Mitosis"]}))
    assert ok(teacher.get(first_path))["hotwords"] == words
    assert ok(admin.get(second_path))["hotwords"] == ["Mitosis"]
    assert teacher.put(second_path, json={"hotwords": ["Wrong course"]}).status_code == 403
    assert teacher.get(second_path).status_code == 403
    assert student.put(first_path, json={"hotwords": []}).status_code == 403
    for invalid in (words + ["extra"], ["x" * 101], [word + "x" for word in words], [12]):
        assert teacher.put(first_path, json={"hotwords": invalid}).status_code == 422
        assert ok(teacher.get(first_path))["hotwords"] == words
    with factory() as db:
        assert speech.exam_policy(db, SimpleNamespace(course_id=first["id"], snapshot={}))['hotwords'] == words
        assert speech.exam_policy(db, SimpleNamespace(course_id=second["id"], snapshot={}))['hotwords'] == ["Mitosis"]
        # An explicitly empty frozen vocabulary must not fall back to current course words.
        assert "hotwords" not in speech.exam_policy(db, SimpleNamespace(course_id=first["id"], snapshot={"hotwords": []}))
    ok(teacher.put(first_path, json={"hotwords": []}))
    assert ok(teacher.get(first_path))["hotwords"] == []
    assert ok(admin.get(second_path))["hotwords"] == ["Mitosis"]
    assert admin.put("/admin/settings/speech", json={"hotwords": ["Global"]}).status_code == 422
    ok(admin.put("/admin/settings/speech", json={"provider": "local_server", "language": "vi"}))


def test_imported_student_google_login_preserves_roster(env, monkeypatch):
    from app import google_login
    from app.models import OAuthFlow
    from app.security import digest

    clients, factory, key, _ = setup_workflow(env)
    office, browser = clients["reviewer"], clients["outsider"]
    ok(import_roster(office, key, [["0009", "google@example.edu.vn", "Google student", ""]]))
    state = "verified-google-import-flow"
    with factory() as db:
        learner = db.scalar(select(User).where(User.email == "google@example.edu.vn"))
        learner_id = learner.id
        flow = OAuthFlow(state_hash=digest(state), poll_hash=digest("test-poll-token-long-enough-123456789"), nonce="nonce", verifier="verifier", expires_at=time.time() + 300)
        db.add(flow)
        db.commit()
        flow_id = flow.id
    monkeypatch.setattr(google_login, "config", lambda: SimpleNamespace(public_origin="http://localhost:3000"))
    monkeypatch.setattr(google_login, "verify_identity", lambda *args: {
        "sub": "google-import-sub", "email": "Google@example.edu.vn", "email_verified": True,
    })
    response = browser.get("/auth/google/callback", params={"state": state, "code": "test"},
                           headers={"Cookie": f"google_oauth_state={state}"}, follow_redirects=False)
    assert response.status_code == 200
    ok(browser.post("/auth/google/poll", json={"flow_id": flow_id, "poll_token": "test-poll-token-long-enough-123456789"}))
    assert ok(browser.get("/auth/me"))["id"] == learner_id
    with factory() as db:
        assert db.get(User, learner_id).google_sub == "google-import-sub"
        assert len(db.scalars(select(User).where(User.email == "google@example.edu.vn")).all()) == 1
        assert db.scalar(select(Assignment).where(Assignment.exam_id == key)).student_id == learner_id


def test_excel_rejects_formulas_and_exports_literal_names(env):
    from openpyxl import Workbook

    clients, _, key, _ = setup_workflow(env)
    office = clients["reviewer"]
    book = Workbook()
    book.active.append(HEADERS)
    book.active.append(["0002", "test@example.edu.vn", '=HYPERLINK("https://example.com")', ""])
    raw = io.BytesIO()
    book.save(raw)
    response = office.post(f"/admin/exams/{key}/import-students", files={"file": ("bad.xlsx", raw.getvalue())})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "EXCEL_FORMULA"
    ok(import_roster(office, key, [["0002", "test@example.edu.vn", "=Literal name", ""]]))
    exported = load_workbook(io.BytesIO(office.get(f"/admin/exams/{key}/grades.xlsx").content))
    assert exported.active.cell(2, 3).value == "=Literal name"
    assert exported.active.cell(2, 3).data_type == "s"
