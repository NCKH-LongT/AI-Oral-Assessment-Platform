import io

from openpyxl import load_workbook
from sqlalchemy import func, select
from test_mvp import ok, prepare

from app import ai
from app.excel import workbook_bytes
from app.excel_imports import SPECS
from app.models import Course, Exam, Rubric, User


def sheet(kind, rows):
    return {"file": ("import.xlsx", workbook_bytes(SPECS[kind][0], rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def test_templates_and_staff_import_permissions_and_idempotency(env):
    clients, factory = env
    admin = clients["admin"]
    for kind in SPECS:
        res = admin.get(f"/admin/imports/{kind}/template.xlsx")
        assert res.status_code == 200
        book = load_workbook(io.BytesIO(res.content))
        assert [c.value for c in book["Data"][1]] == SPECS[kind][0]
        assert "Huong dan" in book.sheetnames
    for kind, role in [("teachers", "TEACHER"), ("officers", "EXAM_OFFICER")]:
        data = [[f"import-{kind}", f"{kind}@example.edu.vn", "Import User", "Unique-password-123"]]
        for actor in ("student", "teacher", "reviewer"):
            assert clients[actor].get(f"/admin/imports/{kind}/template.xlsx").status_code == 403
            assert clients[actor].post(f"/admin/imports/{kind}", files=sheet(kind, data)).status_code == 403
        preview = ok(admin.post(f"/admin/imports/{kind}/preview", files=sheet(kind, data)))
        assert preview["valid"] and preview["create"] == 1
        assert "Unique-password" not in str(preview)
        with factory() as db:
            assert not db.scalar(select(User).where(User.username == data[0][0]))
        assert ok(admin.post(f"/admin/imports/{kind}", files=sheet(kind, data)))["create"] == 1
        assert ok(admin.post(f"/admin/imports/{kind}", files=sheet(kind, data)))["skip"] == 1
        with factory() as db:
            user = db.scalar(select(User).where(User.username == data[0][0]))
            assert user.role == role
    # Existing students/admins cannot be promoted by importing their identities.
    collision = [["student", "student@example.edu.vn", "STUDENT", "Unique-password-123"]]
    assert admin.post("/admin/imports/teachers", files=sheet("teachers", collision)).status_code == 422
    with factory() as db:
        assert db.scalar(select(User).where(User.username == "student")).role == "STUDENT"


def test_atomic_course_import_preview_and_duplicate_protection(env):
    admin = env[0]["admin"]
    rows = [["NEW1", "New course", "Description", ""], ["NEW2", "Other course", "", "missing@example.edu.vn"]]
    preview = ok(admin.post("/admin/imports/courses/preview", files=sheet("courses", rows)))
    assert not preview["valid"] and preview["errors"][0]["row"] == 3
    assert admin.post("/admin/imports/courses", files=sheet("courses", rows)).status_code == 422
    with env[1]() as db:
        assert db.scalar(select(func.count()).select_from(Course)) == 0
    assert ok(admin.post("/admin/imports/courses", files=sheet("courses", rows[:1])))["create"] == 1
    assert ok(admin.post("/admin/imports/courses", files=sheet("courses", rows[:1])))["skip"] == 1
    rows[0][1] = "Overwritten course"
    assert admin.post("/admin/imports/courses", files=sheet("courses", rows[:1])).status_code == 422


def test_rubric_and_exact_questions_import_without_ai_generation(env, monkeypatch):
    context = prepare(env, 1)
    admin = env[0]["admin"]
    code = context["course"]["code"]
    rubric_rows = [[code, "Imported rubric", "Content", "Accurate content", 10, 6], [code, "Imported rubric", "Clarity", "Clear examples", 10, 4]]
    assert ok(admin.post("/admin/imports/rubrics", files=sheet("rubrics", rubric_rows)))["create"] == 1
    assert ok(admin.post("/admin/imports/rubrics", files=sheet("rubrics", rubric_rows)))["skip"] == 1
    rows = [[code, "Imported exam", "Imported rubric", 10, n, context["topic"]["name"], "MEDIUM", text, "dependency; example", "dependency=phụ thuộc"] for n, text in [(2, "Give an example of using dependency injection in a project."), (1, "Explain dependency injection and its main advantages.")]]
    assert ok(admin.post("/admin/imports/exams/preview", files=sheet("exams", rows)))["create"] == 1
    assert ok(admin.post("/admin/imports/exams", files=sheet("exams", rows)))["create"] == 1
    with env[1]() as db:
        exam = db.scalar(select(Exam).where(Exam.name == "Imported exam"))
        key = exam.id
        assert exam.status == "DRAFT" and exam.workflow
        assert exam.imported_questions[0]["text"] == rows[1][7]
        rubric = db.get(Rubric, exam.rubric_id)
        criteria = rubric.criteria
    monkeypatch.setattr(ai, "generate_question", lambda *a, **k: (_ for _ in ()).throw(AssertionError("AI must not rewrite imported questions")))
    monkeypatch.setattr(ai, "structured", lambda *a, **k: (_ for _ in ()).throw(AssertionError("AI must not rewrite imported rubric")))
    result = ok(admin.post(f"/admin/exams/{key}/generate"))
    assert [q["text"] for q in result["snapshot"]["questions"]] == [rows[1][7], rows[0][7]]
    assert result["snapshot"]["criteria"] == criteria
    assert all(q["reference_chunk_ids"] for q in result["snapshot"]["questions"])
    assert admin.post("/admin/imports/exams", files=sheet("exams", rows)).status_code == 422
    assert env[0]["teacher"].post("/admin/imports/exams/preview", files=sheet("exams", rows)).status_code == 403


def test_import_rejects_formulas_bad_headers_and_duplicate_question_numbers(env):
    admin = env[0]["admin"]
    data = workbook_bytes(SPECS["courses"][0], [["CODE", "Course", "", ""]])
    book = load_workbook(io.BytesIO(data))
    book.active.cell(2, 2).value = '=HYPERLINK("https://example.com")'
    out = io.BytesIO()
    book.save(out)
    response = admin.post("/admin/imports/courses/preview", files={"file": ("test.xlsx", out.getvalue())})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "EXCEL_FORMULA"
    assert admin.post("/admin/imports/courses/preview", files={"file": ("test.xlsx", workbook_bytes(["wrong"], [[1]]))}).status_code == 422
    ctx = prepare(env, 1)
    row = [ctx["course"]["code"], "Duplicate questions", ctx["rubric"]["name"], 10, 1, ctx["topic"]["name"], "EASY", "Explain dependency injection in your own words.", "definition", ""]
    result = ok(admin.post("/admin/imports/exams/preview", files=sheet("exams", [row, row])))
    assert not result["valid"] and "Số câu" in result["errors"][0]["message"]
