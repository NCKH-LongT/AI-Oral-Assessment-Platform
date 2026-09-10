import hashlib
import time

import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError
from sqlalchemy import select

from app import ai, worker
from app.documents import chunk_text, extract
from app.main import app
from app.models import Attempt, Exam, ExamSession
from app.schemas import GradeOutput, RubricIn


def ok(response, status=200):
    assert response.status_code == status, response.text
    return response.json()


def prepare(env, count=2):
    clients, factory = env
    admin = clients["admin"]
    course = ok(admin.post("/admin/courses", json={"code": "SE101", "name": "Software Engineering"}), 201)
    root = f"/admin/courses/{course['id']}"
    lo = ok(admin.post(root + "/outcomes", json={"code": "LO1", "description": "Explain DI"}), 201)
    topic = ok(
        admin.post(root + "/topics", json={"name": "Dependency Injection", "learning_outcome_id": lo["id"]}),
        201,
    )
    doc = ok(
        admin.post(
            root + "/documents",
            data={"topic_id": topic["id"]},
            files={
                "file": (
                    "guide.txt",
                    b"Dependency injection supplies dependencies externally. Constructor injection improves testing.",
                    "text/plain",
                )
            },
        ),
        201,
    )
    assert worker.tick()
    assert ok(admin.get(root + "/workspace"))["documents"][0]["status"] == "READY"
    rubric = ok(
        admin.post(
            root + "/rubrics",
            json={
                "name": "Rubric",
                "criteria": [
                    {"name": "knowledge", "description": "Correctness", "max_score": 5, "weight": 3},
                    {"name": "explanation", "description": "Clear example", "max_score": 5, "weight": 2},
                ],
            },
        ),
        201,
    )
    payload = {
        "course_id": course["id"],
        "rubric_id": rubric["id"],
        "name": "Oral test",
        "time_limit": 900,
        "blueprint": [{"topic_id": topic["id"], "difficulty": "MEDIUM", "count": count}],
    }
    exam = ok(admin.post("/admin/exams", json=payload), 201)
    ok(admin.post(f"/admin/exams/{exam['id']}/publish"))
    student_id = ok(clients["student"].get("/auth/me"))["id"]
    ok(admin.post(f"/admin/exams/{exam['id']}/assign", json={"student_ids": [student_id]}))
    return {
        "course": course,
        "topic": topic,
        "document": doc,
        "rubric": rubric,
        "exam": exam,
        "payload": payload,
    }


def upload(client, attempt_id, kind, corrupt=False):
    blob = b"\x1aE\xdf\xa3" + b"test-container-bytes" * 40
    checksum = hashlib.sha256(blob).hexdigest()
    data = {
        "attempt_id": attempt_id,
        "kind": kind,
        "mime_type": f"{kind.lower()}/webm",
        "size": len(blob),
        "sha256": checksum,
    }
    row = ok(client.post("/uploads/init", json=data))
    assert client.post(f"/uploads/{row['id']}/complete").status_code == 409
    if corrupt:
        assert (
            client.put(
                f"/uploads/{row['id']}/chunks/0", content=blob, headers={"X-Chunk-Sha256": "0" * 64}
            ).status_code
            == 422
        )
    ok(client.put(f"/uploads/{row['id']}/chunks/0", content=blob, headers={"X-Chunk-Sha256": checksum}))
    ok(client.post(f"/uploads/{row['id']}/complete"))
    ok(client.post(f"/uploads/{row['id']}/complete"))
    return row, blob


def start(env, context):
    student = env[0]["student"]
    session = ok(student.post("/exam-sessions", json={"exam_id": context["exam"]["id"]}))
    assert session["current_attempt"] is None
    return ok(student.post(f"/exam-sessions/{session['id']}/start"))


def test_end_to_end_demo_with_evidence_and_review(env):
    context = prepare(env)
    clients, factory = env
    student, admin = clients["student"], clients["admin"]
    exams = ok(student.get("/exams/available"))
    assert len(exams) == 1
    session = start(env, context)
    sid = session["id"]
    assert student.post(f"/exam-sessions/{sid}/finish").status_code == 409
    assert "expected_concepts" not in str(session) and "criteria" not in str(session)
    for _ in range(2):
        attempt = session["current_attempt"]
        path = f"/question-attempts/{attempt['id']}"
        ok(student.post(path + "/start"))
        body = {
            "transcript": "Dependency injection provides dependencies externally.",
            "stt_confidence": 0.96,
        }
        ok(student.post(path + "/submit", json=body, headers={"Idempotency-Key": "test-submit-123"}))
        ok(student.post(path + "/submit", json=body, headers={"Idempotency-Key": "test-submit-123"}))
        assert (
            student.post(
                path + "/submit", json=body, headers={"Idempotency-Key": "different-key-123"}
            ).status_code
            == 409
        )
        video, blob = upload(student, attempt["id"], "VIDEO", corrupt=True)
        upload(student, attempt["id"], "AUDIO")
        assert admin.get(f"/evidence/{video['id']}/content").content == blob
        partial = admin.get(f"/evidence/{video['id']}/content", headers={"Range": "bytes=4-12"})
        assert partial.status_code == 206 and partial.content == blob[4:13]
        assert clients["outsider"].get(f"/evidence/{video['id']}/content").status_code == 403
        assert worker.tick()
        session = ok(student.get(f"/exam-sessions/{sid}"))
    final = ok(student.post(f"/exam-sessions/{sid}/finish"))
    assert final["status"] == "REVIEW_REQUIRED" and final["final_score"] is None
    assert ok(student.post(f"/exam-sessions/{sid}/finish")) == final or final["status"] == "REVIEW_REQUIRED"
    review = ok(admin.get(f"/admin/results/{sid}"))
    assert len(review["attempts"]) == 2
    for a in review["attempts"]:
        assert a["assessment"]["retrieved_chunks"]
        assert a["assessment"]["rubric_version"] == 1
        assert len(a["evidence"]) == 2


def test_authorization_assignment_snapshot_and_rag_isolation(env):
    context = prepare(env)
    clients, factory = env
    assert ok(clients["outsider"].get("/exams/available")) == []
    assert (
        clients["outsider"].post("/exam-sessions", json={"exam_id": context["exam"]["id"]}).status_code == 403
    )
    assert clients["student"].get("/admin/courses").status_code == 403
    assert clients["reviewer"].post("/admin/courses", json={"code": "x", "name": "x"}).status_code == 403
    assert clients["teacher"].get(f"/admin/courses/{context['course']['id']}/workspace").status_code == 403
    assert (
        clients["admin"].put(f"/admin/exams/{context['exam']['id']}", json=context["payload"]).status_code
        == 409
    )
    session = start(env, context)
    assert clients["outsider"].get(f"/exam-sessions/{session['id']}").status_code == 403
    new_rubric = {
        "name": "Updated",
        "criteria": [{"name": "new", "description": "new", "max_score": 10, "weight": 1}],
    }
    ok(clients["admin"].put(f"/admin/rubrics/{context['rubric']['id']}", json=new_rubric))
    with factory() as db:
        snapshot = db.get(Exam, context["exam"]["id"]).snapshot
        assert snapshot["rubric_version"] == 1 and snapshot["criteria"][0]["name"] == "knowledge"
        assert ai.retrieve(db, "different-course", context["topic"]["id"], "DI") == []
        assert ai.retrieve(db, context["course"]["id"], context["topic"]["id"], "DI", []) == []


def test_state_order_deadline_and_missing_evidence(env):
    context = prepare(env)
    clients, factory = env
    session = start(env, context)
    student = clients["student"]
    with factory() as db:
        attempts = db.scalars(
            select(Attempt).where(Attempt.session_id == session["id"]).order_by(Attempt.sequence)
        ).all()
        first, second = [a.id for a in attempts]
    assert student.post(f"/question-attempts/{second}/start").status_code == 409
    assert (
        student.post(
            f"/question-attempts/{first}/submit",
            json={"transcript": "answer", "stt_confidence": 1},
            headers={"Idempotency-Key": "no-start-123"},
        ).status_code
        == 409
    )
    ok(student.post(f"/question-attempts/{first}/start"))
    with factory() as db:
        db.get(ExamSession, session["id"]).started_at = time.time() - 1000
        db.commit()
    body = {"transcript": "Late but preserved", "stt_confidence": 0.99}
    ok(
        student.post(
            f"/question-attempts/{first}/submit", json=body, headers={"Idempotency-Key": "late-answer-123"}
        )
    )
    ok(
        student.post(
            f"/question-attempts/{first}/submit", json=body, headers={"Idempotency-Key": "late-answer-123"}
        )
    )
    assert student.post(f"/question-attempts/{second}/start").status_code == 409
    assert student.post(f"/exam-sessions/{session['id']}/finish").status_code == 409
    upload(student, first, "AUDIO")
    upload(student, first, "VIDEO")
    ok(student.post(f"/exam-sessions/{session['id']}/finish"))
    worker.tick()
    assert ok(student.get(f"/exam-sessions/{session['id']}"))["status"] == "REVIEW_REQUIRED"


def test_auth_rotation_logout_csrf_errors(env):
    client = env[0]["student"]
    old_cookies = dict(client.cookies)
    ok(client.post("/auth/refresh"))
    with TestClient(app) as other:
        other.cookies.update(old_cookies)
        assert other.get("/auth/me").status_code == 401
        assert other.post("/auth/refresh").status_code == 401
    assert client.post("/auth/logout", headers={"Origin": "https://evil.invalid"}).status_code == 403
    ok(client.post("/auth/logout"))
    assert client.get("/auth/me").status_code == 401
    bad = client.post("/auth/login", json={"username": "admin", "password": "wrong"})
    assert bad.status_code == 401 and bad.json()["error"]["code"] == "INVALID_CREDENTIALS"
    invalid = client.post("/auth/login", json={"username": "", "password": "DO-NOT-ECHO", "extra": 2})
    assert invalid.status_code == 422 and "DO-NOT-ECHO" not in invalid.text


def test_grading_failure_keeps_transcript_for_review(env, monkeypatch):
    context = prepare(env, 1)
    session = start(env, context)
    student = env[0]["student"]
    aid = session["current_attempt"]["id"]
    ok(student.post(f"/question-attempts/{aid}/start"))
    ok(
        student.post(
            f"/question-attempts/{aid}/submit",
            json={"transcript": "saved answer", "stt_confidence": 1},
            headers={"Idempotency-Key": "safe-transcript"},
        )
    )

    def broken(*args, **kwargs):
        raise TimeoutError("upstream timeout")

    monkeypatch.setattr(ai, "grade", broken)
    assert worker.tick()
    with env[1]() as db:
        attempt = db.get(Attempt, aid)
        assert attempt.transcript == "saved answer"
        assert attempt.assessment["review_required"] and attempt.assessment["score"] is None


def test_document_failure_and_retry(env):
    context = prepare(env)
    admin = env[0]["admin"]
    doc = ok(
        admin.post(
            f"/admin/courses/{context['course']['id']}/documents",
            data={"topic_id": context["topic"]["id"]},
            files={"file": ("bad.pdf", b"not a pdf", "application/pdf")},
        ),
        201,
    )
    worker.tick()
    workspace = ok(admin.get(f"/admin/courses/{context['course']['id']}/workspace"))
    assert next(d for d in workspace["documents"] if d["id"] == doc["id"])["status"] == "FAILED"
    ok(admin.post(f"/admin/documents/{doc['id']}/retry"))
    assert worker.tick()


def test_structured_grading_server_calculation_and_validation(monkeypatch):
    cfg = ai.settings()
    monkeypatch.setattr(cfg, "ai_provider", "gemini")
    criteria = [
        {"name": "a", "description": "a", "max_score": 5, "weight": 3},
        {"name": "b", "description": "b", "max_score": 5, "weight": 2},
    ]
    value = {
        "confidence": 0.95,
        "criteria": [{"name": "a", "score": 4, "comment": "ok"}, {"name": "b", "score": 3, "comment": "ok"}],
        "missing_concepts": [],
        "reasoning_summary": "Grounded",
        "reference_chunk_ids": ["chunk-1"],
    }
    monkeypatch.setattr(ai, "structured", lambda *args: GradeOutput.model_validate(value))
    result = ai.grade({}, "answer", criteria, [{"id": "chunk-1"}], 0.98)
    assert result["score"] == 7.2 and not result["review_required"]
    assert ai.grade({}, "answer", criteria, [{"id": "chunk-1"}], 0.1)["review_required"]
    value["criteria"][0]["score"] = 6
    with pytest.raises(ValueError, match="exceeds"):
        ai.grade({}, "answer", criteria, [{"id": "chunk-1"}], 0.99)
    value["criteria"][0]["score"] = 4
    value["reference_chunk_ids"] = ["hallucinated"]
    with pytest.raises(ValueError, match="references"):
        ai.grade({}, "answer", criteria, [{"id": "chunk-1"}], 0.99)


def test_chunking_and_rubric_validation():
    chunks = list(chunk_text("first paragraph\n\nsecond paragraph", 20))
    assert chunks == ["first paragraph", "second paragraph"]
    assert all(len(c) <= 100 for c in chunk_text("x" * 350, 100))
    assert extract(b"valid utf8", "sample.txt") == [(1, "valid utf8")]
    with pytest.raises(ValueError):
        extract(b"bad", "file.exe")
    with pytest.raises(ValidationError):
        RubricIn.model_validate(
            {"name": "bad", "criteria": [{"name": "same", "description": "x", "max_score": 5}] * 2}
        )


def test_office_extraction_preserves_page_and_content():
    import io

    from docx import Document
    from pptx import Presentation

    word = Document()
    word.add_paragraph("Dependency injection")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Constructor"
    table.cell(0, 1).text = "External dependency"
    buffer = io.BytesIO()
    word.save(buffer)
    content = extract(buffer.getvalue(), "guide.docx")
    assert content[0][0] == 1 and "External dependency" in content[0][1]
    deck = Presentation()
    for i in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"Topic {i + 1}"
        slide.placeholders[1].text = "Course knowledge"
    buffer = io.BytesIO()
    deck.save(buffer)
    content = extract(buffer.getvalue(), "guide.pptx")
    assert [page for page, _ in content] == [1, 2]
    assert "Topic 2" in content[1][1]
