import io
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from test_mvp import ok, prepare, start, upload

from app import storage, worker
from app.config import settings
from app.models import Attempt, Course, Exam, ExamSession, MediaCleanup, ReviewJob


def completed(env):
    context = prepare(env, 1)
    session = start(env, context)
    aid = session["current_attempt"]["id"]
    student = env[0]["student"]
    ok(student.post(f"/question-attempts/{aid}/start"))
    evidence, blob = upload(student, aid, "SCREEN")
    with env[1]() as db:
        sitting = db.get(ExamSession, session["id"])
        sitting.status, sitting.ai_score, sitting.final_score = "COMPLETED", 8, 8
        attempt = db.get(Attempt, aid)
        attempt.status, attempt.transcript, attempt.assessment = "GRADED", "Original answer", {"score": 8}
        db.commit()
    return context, session, evidence, blob


def preview(admin, sid, mode="MEDIA"):
    return ok(admin.post("/admin/storage/preview", json={"session_ids": [sid], "mode": mode}))


def confirm(admin, plan):
    return admin.post("/admin/storage/cleanup", json={"token": plan["token"], "confirmation": plan["confirmation"], "reason": "Retention period ended"})


def test_report_filter_export_and_media_only_cleanup_preserves_scores(env):
    ctx, session, evidence, blob = completed(env)
    admin = env[0]["admin"]
    report = ok(admin.get("/admin/storage"))
    assert report["total"] == 1 and report["media_bytes"] == len(blob) and report["files"] == 1
    assert report["rows"][0]["media_bytes"] == len(blob)
    assert report["by_course"][0]["id"] == ctx["course"]["id"]
    assert report["by_student"][0]["bytes"] == len(blob)
    assert ok(admin.get("/admin/storage?course_id=unmatched"))["media_bytes"] == 0
    assert ok(admin.get("/admin/storage?student_id=unmatched"))["total"] == 0
    book = load_workbook(io.BytesIO(admin.get("/admin/storage/report.xlsx").content))
    assert book.active.cell(2, 9).value == len(blob)
    plan = preview(admin, session["id"])
    assert plan["media_bytes"] == len(blob)
    # A mistyped confirmation cannot delete anything.
    assert admin.post("/admin/storage/cleanup", json={"token": plan["token"], "confirmation": "XOA 999", "reason": "Retention ended"}).status_code == 422
    assert env[0]["student"].get(f"/evidence/{evidence['id']}/content").status_code == 200
    ok(confirm(admin, plan))
    assert confirm(admin, plan).status_code == 409
    detail = ok(admin.get(f"/admin/results/{session['id']}"))
    assert detail["final_score"] == 8 and detail["media_removed_at"]
    assert detail["attempts"][0]["transcript"] == "Original answer"
    assert detail["attempts"][0]["evidence"] == []
    assert admin.get(f"/evidence/{evidence['id']}/content").status_code == 404
    report = ok(admin.get("/admin/storage"))
    assert report["media_bytes"] == 0 and report["cleanup"]["estimated_bytes"] == len(blob)
    assert storage.get(f"evidence/{evidence['id']}") == blob  # deletion waits for worker commit
    assert worker.tick()
    assert ok(admin.get("/admin/storage"))["cleanup"]["files"] == 0
    assert not storage.local_path(f"evidence/{evidence['id']}").exists()


def test_full_cleanup_keeps_course_exam_account_and_unselected_session(env):
    ctx, session, _, _ = completed(env)
    clients, factory = env
    other = ok(clients["outsider"].get("/auth/me"))["id"]
    ok(clients["admin"].post(f"/admin/exams/{ctx['exam']['id']}/assign", json={"student_ids": [other]}))
    untouched = ok(clients["outsider"].post("/exam-sessions", json={"exam_id": ctx["exam"]["id"]}))
    plan = preview(clients["admin"], session["id"], "ALL")
    ok(confirm(clients["admin"], plan))
    with factory() as db:
        assert db.get(Course, ctx["course"]["id"]) and db.get(Exam, ctx["exam"]["id"])
        deleted = db.get(ExamSession, session["id"])
        assert deleted.status == "DELETED" and deleted.final_score is None
        assert not db.scalar(select(Attempt.id).where(Attempt.session_id == session["id"]))
        assert db.get(ExamSession, untouched["id"]).status == "DEVICE_CHECK"
    assert clients["admin"].get(f"/admin/results/{session['id']}").status_code == 404
    assert ok(clients["admin"].get("/admin/storage"))["total"] == 1


def test_cleanup_rejects_stale_selection_active_exams_and_pending_review(env):
    _, session, evidence, _ = completed(env)
    admin, factory = env[0]["admin"], env[1]
    plan = preview(admin, session["id"])
    with factory() as db:
        db.get(ExamSession, session["id"]).final_score = 7
        db.commit()
    assert confirm(admin, plan).status_code == 409
    assert admin.get(f"/evidence/{evidence['id']}/content").status_code == 200
    with factory() as db:
        db.get(ExamSession, session["id"]).status = "IN_PROGRESS"
        db.commit()
    assert admin.post("/admin/storage/preview", json={"session_ids": [session["id"]]}).status_code == 409
    assert confirm(admin, plan).status_code == 409
    with factory() as db:
        db.get(ExamSession, session["id"]).status = "COMPLETED"
        attempt = db.scalar(select(Attempt).where(Attempt.session_id == session["id"]))
        db.add(ReviewJob(attempt_id=attempt.id, requested_by=ok(admin.get("/auth/me"))["id"], reason="Recheck grading", policy={}, original={}))
        db.commit()
    assert admin.post("/admin/storage/preview", json={"session_ids": [session["id"]]}).status_code == 409


@pytest.mark.parametrize("role", ["student", "teacher", "reviewer"])
def test_storage_is_admin_only(env, role):
    client = env[0][role]
    for path in ("/admin/storage", "/admin/storage/report.xlsx"):
        assert client.get(path).status_code == 403
    assert client.post("/admin/storage/preview", json={"session_ids": ["synthetic"]}).status_code == 403
    assert client.post("/admin/storage/cleanup", json={"token": "fake", "confirmation": "XOA 1", "reason": "Cleanup requested"}).status_code == 403


def test_failed_object_deletion_is_reported_and_retried_including_staging(env, monkeypatch):
    _, session, evidence, blob = completed(env)
    admin = env[0]["admin"]
    folder = Path(settings().data_dir) / "uploads" / evidence["id"]
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "leftover").write_bytes(b"staging")
    plan = preview(admin, session["id"])
    assert plan["staging_bytes"] == 7
    ok(confirm(admin, plan))
    original = storage.delete
    monkeypatch.setattr(storage, "delete", lambda _: (_ for _ in ()).throw(OSError("storage down")))
    assert worker.tick()
    report = ok(admin.get("/admin/storage"))
    assert report["cleanup"]["retrying"] == 1
    assert report["cleanup"]["estimated_bytes"] == len(blob) + 7
    with env[1]() as db:
        db.scalar(select(MediaCleanup)).next_attempt_at = time.time() - 1
        db.commit()
    monkeypatch.setattr(storage, "delete", original)
    assert worker.tick()
    assert not folder.exists()
    assert ok(admin.get("/admin/storage"))["cleanup"]["files"] == 0
